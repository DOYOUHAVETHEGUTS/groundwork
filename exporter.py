"""Exports: Word (five literal sections + appendix), Markdown, and a JSON payload for Levelpath."""
import io
import json
import re

import rubric
import store

DRAFT_NOTE = "DRAFT — not approval-ready. Groundwork score {score}/{th}{gates}."


def _f(v):
    return str(v or "").strip()


def _money(v):
    return f"${v:,.2f}" if v % 1 else f"${v:,.0f}"


def _cost_table(req):
    items = [i for i in req.get("cost_items") or [] if _f(i.get("desc")) and rubric._num(i.get("unit")) > 0]
    if not items:
        return None
    rows = [(i["desc"], f"{rubric._num(i.get('qty') or 1):g}", _money(rubric._num(i["unit"])),
             _money(rubric._num(i.get("qty") or 1) * rubric._num(i["unit"]))) for i in items]
    sub = sum(rubric._num(i.get("qty") or 1) * rubric._num(i["unit"]) for i in items)
    pct = rubric._num(req.get("contingency_pct"))
    if pct:
        rows.append((f"Contingency ({pct:g}%)", "", "", _money(sub * pct / 100)))
    rows.append(("Total", "", "", _money(sub * (1 + pct / 100))))
    return ["Description", "Qty", "Unit cost", "Amount"], rows


def _quote_summary(req):
    vq = rubric.valid_quotes(req)
    if not vq:
        return None
    rows = [(q["vendor"] + (" (selected)" if q.get("selected") else ""), _money(rubric.quote_total(q)),
             _f(q.get("lead_time")) or "—", _f(q.get("warranty")) or "—",
             str(len(rubric.quote_items(q))) if rubric.quote_items(q) else "not itemized") for q in vq]
    return ["Vendor", "Total", "Lead time", "Warranty", "Line items"], rows


def _narrative_cost(req):
    t = _f(req.get("cost"))
    if req.get("cost_items"):
        t = "\n".join(l for l in t.split("\n") if not (l.strip().startswith("-") and "×" in l))
    return t.strip()


def _status_line(req):
    q = req.get("qualification") or {}
    if not q:
        return ""
    gates = [g["label"] for g in q.get("gates") or [] if not g.get("passed")]
    if q.get("overall", 0) >= q.get("threshold", 75) and not gates:
        return ""
    return DRAFT_NOTE.format(score=q.get("overall"), th=q.get("threshold", 75),
                             gates=("; required items open: " + ", ".join(gates)) if gates else "")


def coaching_summary(req):
    co = req.get("coaching") or {}
    rounds = [r for r in co.get("rounds") or [] if r.get("answered")]
    if not rounds:
        return None
    esc = co.get("escalation") or {}
    return {"status": co.get("status"), "original_score": co.get("original_score"),
            "rounds": [{"round": r["round"], "score_before": r.get("score_before"), "score_after": r.get("score_after"),
                        "passed": r.get("passed"),
                        "questions": [{"question": x["question"], "field": x["field"],
                                       "answer": (r.get("answers") or {}).get(x["id"], "")} for x in r.get("questions", [])],
                        "largest_edits": [{"field": e["label"], "points": e["points"]} for e in (r.get("edits") or [])[:5]]}
                       for r in rounds],
            "escalation": {"to": esc.get("to_email"), "method": esc.get("method"), "delivered": esc.get("delivered"),
                           "simulated": esc.get("simulated")} if esc else None}


def levelpath_payload(req):
    q = req.get("qualification") or {}
    return {
        "schema": "groundwork.five_point.v2", "request_id": req.get("id"), "title": req.get("title"),
        "department": req.get("department"), "request_type": req.get("request_type"), "requester": req.get("requester"),
        "five_point": {k: req.get(k, "") for k, _ in rubric.SECTIONS}, "appendix": req.get("appendix", ""),
        "cre": {"driven": req.get("cre_driven"), "template": req.get("cre_template")},
        "cost_items": req.get("cost_items") or [], "contingency_pct": req.get("contingency_pct"),
        "funding_source": req.get("funding_source"),
        "quotes": [{k: x.get(k) for k in ("vendor", "date", "reference", "items", "lead_time", "warranty", "notes", "selected")}
                   | {"total": rubric.quote_total(x)} for x in rubric.valid_quotes(req)],
        "selection_rationale": req.get("selection_rationale", ""),
        "exhibits": [{k: e.get(k) for k in ("label", "title", "date", "kind", "citation")} for e in req.get("exhibits") or []],
        "qualification": {"score": q.get("overall"), "status": q.get("status"),
                          "sections": [{"name": c["name"], "earned": c.get("earned"), "max": c["weight"],
                                        "findings": [f["title"] for f in c.get("findings", [])]} for c in q.get("categories", [])],
                          "gates": q.get("gates"), "engine": q.get("engine")},
        "guided_improvement": coaching_summary(req),
    }


def to_json_bytes(req):
    return (json.dumps(levelpath_payload(req), indent=2).encode(), f"{req.get('id', 'request')}-levelpath.json",
            "application/json")


def _md_table(head, rows):
    return ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)] + ["| " + " | ".join(map(str, r)) + " |" for r in rows]


def to_markdown(req):
    L = [f"# {req.get('title') or 'Capital Request'}",
         f"*{req.get('department', '')} · {req.get('request_type', '')} · {req.get('id', '')}*", ""]
    if _status_line(req):
        L += [f"> **{_status_line(req)}**", ""]
    if req.get("cre_driven") == "yes":
        L += [f"**CRE-driven request.** CRE template: {req.get('cre_template') or 'not yet requested'}.", ""]
    for n, (k, lbl) in enumerate(rubric.SECTIONS, 1):
        L += [f"## {n}. {lbl}", (_narrative_cost(req) if k == "cost" else _f(req.get(k))) or "_Not provided_", ""]
        if k == "cost" and _cost_table(req):
            L += _md_table(*_cost_table(req)) + [""]
            if req.get("funding_source"):
                L += [f"**Funding source:** {req['funding_source']}", ""]
        if k == "alternatives":
            if _quote_summary(req):
                L += ["**Vendor quotes**", ""] + _md_table(*_quote_summary(req)) + [""]
            if req.get("selection_rationale"):
                L += [f"**Selected vendor and why:** {req['selection_rationale']}", ""]
    L += ["## Appendix", _f(req.get("appendix")), ""]
    for e in req.get("exhibits") or []:
        L.append(f"- **{e.get('label')}** — {e.get('title')}" + (f" ({e['date']})" if e.get("date") else "")
                 + (f". Source: {e['citation']}" if e.get("citation") else ""))
    return "\n".join(L).strip() + "\n"


def to_docx_bytes(req):
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        return to_markdown(req).encode(), f"{req.get('id', 'request')}-5point.md", "text/markdown"
    import images

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    doc.add_heading(req.get("title") or "Capital Request", 0)
    meta = doc.add_paragraph(" | ".join(x for x in (req.get("department"), req.get("request_type"), req.get("id"),
                                                     req.get("requester")) if x))
    meta.runs[0].font.size = Pt(9)
    if _status_line(req):
        p = doc.add_paragraph()
        r = p.add_run(_status_line(req))
        r.bold, r.font.color.rgb = True, RGBColor(0xC6, 0x46, 0x2E)
    if req.get("cre_driven") == "yes":
        doc.add_paragraph(f"CRE-driven request. CRE template: {req.get('cre_template') or 'not yet requested'}.")

    def table(head, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(head))
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(head):
            t.rows[0].cells[i].text = h
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                if i and re.match(r"^[\d$,.%() -]+$", str(v)):
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph()

    def text(t):
        for line in _f(t).split("\n"):
            line = line.strip()
            if not line:
                continue
            if "|" in line:
                doc.add_paragraph(line.replace("|", "  |  ")).runs[0].font.size = Pt(9)
            elif line.startswith(("-", "•", "*")):
                doc.add_paragraph(line.lstrip("-•* ").strip(), style="List Bullet")
            else:
                doc.add_paragraph(line)

    def picture(png, caption):
        doc.add_picture(io.BytesIO(png), width=Inches(6.2))
        c = doc.add_paragraph(caption)
        c.runs[0].italic, c.runs[0].font.size = True, Pt(8.5)

    exhibits = req.get("exhibits") or []
    for n, (k, lbl) in enumerate(rubric.SECTIONS, 1):
        doc.add_heading(f"{n}. {lbl}", 1)
        text(_narrative_cost(req) if k == "cost" else req.get(k) or "Not provided.")
        if k == "current_situation":
            if len([s for s in req.get("flow_steps") or [] if _f(s.get("name"))]) >= 2:
                try:
                    picture(images.flow_png(req["flow_steps"], "Process flow"), "Process flow prepared from requester input. Red marks the bottleneck.")
                except Exception:
                    pass
            for e in [e for e in exhibits if e.get("image_id") and e.get("section", "current_situation") == "current_situation"][:3]:
                im = store.get_image(e["image_id"])
                if im:
                    picture(im["data"], f"{e.get('label')}: {e.get('title')}" + (f". Source: {e['citation']}" if e.get("citation") else ""))
        if k == "cost":
            if _cost_table(req):
                table(*_cost_table(req))
            if req.get("funding_source"):
                doc.add_paragraph(f"Funding source: {req['funding_source']}")
        if k == "alternatives":
            if _quote_summary(req):
                doc.add_paragraph().add_run("Vendor quotes").bold = True
                table(*_quote_summary(req))
            if req.get("selection_rationale"):
                p = doc.add_paragraph()
                p.add_run("Selected vendor and why: ").bold = True
                p.add_run(req["selection_rationale"])

    doc.add_heading("Appendix & Supporting Detail", 1)
    text(req.get("appendix"))
    for q in rubric.valid_quotes(req):
        if rubric.quote_items(q):
            doc.add_heading(f"Quote detail — {q['vendor']}" + (f" ({q['date']})" if q.get("date") else ""), 2)
            table(["Line item", "Qty", "Unit", "Amount"],
                  [(i["desc"], f"{rubric._num(i.get('qty') or 1):g}", _money(rubric._num(i["unit"])),
                    _money(rubric._num(i.get("qty") or 1) * rubric._num(i["unit"]))) for i in rubric.quote_items(q)]
                  + [("Total", "", "", _money(rubric.quote_total(q)))])
    for e in exhibits:
        if e.get("image_id") and e.get("section", "current_situation") != "current_situation":
            im = store.get_image(e["image_id"])
            if im:
                picture(im["data"], f"{e.get('label')}: {e.get('title')}" + (f". Source: {e['citation']}" if e.get("citation") else ""))
        elif not e.get("image_id"):
            doc.add_paragraph(f"{e.get('label')}: {e.get('title')}" + (f" ({e['date']})" if e.get("date") else ""),
                              style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return (buf.getvalue(), f"{req.get('id', 'request')}-5point.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
